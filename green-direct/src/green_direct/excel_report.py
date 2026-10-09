"""成果输出工作簿（v4.2）：单一 Excel 交付物，多 Sheet + 全中文释义。

设计约定：
- 工作簿是面向用户的唯一交付物；CSV 保留为机器通道（校验器/工作台），
  列名与结构不变。
- 按项目组件过滤：禁用的组件（柴油/生物质）不出现在任何 Sheet、任何
  标签中——风光储项目的工作簿里没有柴/生字样。
- 两层中文释义：①指标/列名固定中文映射（本模块为唯一来源）；
  ②总览每行附"说明"列解释统计口径。
- 任何写簿失败都静默降级为仅 CSV（不阻断求解、不透传内部报错）。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

PROJECT_TYPE_LABELS = {"grid_connected": "并网型", "offgrid": "离网型"}
SOLVER_MODE_LABELS = {"primary": "天枢", "alternate": "天璇", "dual": "双擎互证"}

# 逐时结果列的固定中文映射（机器列名 → 用户列名）。
HOURLY_COLUMN_LABELS: dict[str, str] = {
    "hour_index": "小时序号",
    "load": "负荷功率",
    "buy_price": "购电价",
    "sell_price": "售电价",
    "wind_profile": "风电出力系数",
    "pv_profile": "光伏出力系数",
    "wind_generation": "风电出力",
    "pv_generation": "光伏出力",
    "green_generation": "绿电出力（风+光）",
    "green_to_load": "绿电直供负荷",
    "green_to_storage": "绿电充储",
    "green_to_sell": "绿电上网",
    "green_curtail": "绿电弃电",
    "grid_to_load": "购电功率",
    "unmet_load": "失负荷功率",
    "storage_to_load": "储能放电",
    "soc_energy": "储能电量",
    "soc_ratio": "储能SOC",
    "diesel_to_load": "柴油出力",
    "biomass_generation": "生物质出力",
    "biomass_to_load": "生物质直供",
    "biomass_to_sell": "生物质上网",
    "biomass_curtail": "生物质弃电",
    "biomass_maintenance": "检修标记",
    "biomass_balance_check": "生物质平衡校核",
    "green_balance_check": "绿电平衡校核",
    "load_balance_check": "负荷平衡校核",
}

_HEADER_FONT = Font(bold=True)


def _round(value: float, digits: int = 4) -> float:
    return round(float(value), digits)


def _add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]],
               widths: list[int] | None = None):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
    for row in rows:
        sheet.append(row)
    for index in range(1, len(headers) + 1):
        width = widths[index - 1] if widths and index <= len(widths) else 22
        sheet.column_dimensions[get_column_letter(index)].width = width
    return sheet


def write_results_workbook(
    excel_path: str | Path,
    *,
    summary: dict[str, Any],
    hourly_rows: list[dict[str, Any]],
    project_type: str,
    diesel_enabled: bool,
    biomass_enabled: bool,
    biomass_continuous: bool,
    maintenance_hours: int,
    maintenance_start_hour: int,
    component_bounds: dict[str, tuple[float, float]] | None = None,
    boundary_notes: list[str] | tuple[str, ...] = (),
    solver_mode: str | None = None,
) -> Path:
    """生成多 Sheet 中文工作簿；调用方负责异常兜底。

    component_bounds：各组件 (下限, 上限)，键为 wind/pv/ess_energy/ess_power/
    diesel/biomass/ biomass_energy（生物质年发电量上限），仅启用的组件传入。
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    bounds = component_bounds or {}
    grid = project_type == "grid_connected"

    # ── 方案总览 ──
    overview_rows: list[list[Any]] = [
        ["项目类型", PROJECT_TYPE_LABELS.get(project_type, project_type), "并网型可与电网购售电，离网型自平衡"],
        ["全生命周期总成本（NPV）", _round(float(summary["objective_value"]), 2), "元；全生命周期折现总成本，越小越优"],
        ["初始投资", _round(float(summary["initial_investment_cost"]), 2), "元；一次性建设投入"],
        ["年运维", _round(float(summary["annual_om_cost"]), 2), "元/年"],
    ]
    if grid:
        overview_rows.append(["年购电成本", _round(float(summary["annual_buy_cost"]), 2), "元/年"])
        overview_rows.append(["年售电收益", _round(float(summary["annual_sell_revenue"]), 2), "元/年"])
    else:
        overview_rows.append(["年失负荷惩罚成本", _round(float(summary["annual_unmet_penalty_cost"]), 2), "元/年；停电电量的惩罚计价"])
    overview_rows += [
        ["绿电自用率", _round(float(summary["self_use_ratio_of_green"]), 6),
         "自用电量 / 总发电量" + ("（含生物质）" if biomass_enabled else "")],
        ["弃电率", _round(float(summary["curtail_ratio_of_green"]), 6), "弃电量 / 总发电量"],
    ]
    if grid:
        overview_rows.append(["绿电外售率", _round(float(summary["sell_ratio_of_green"]), 6), "上网电量 / 总发电量"])
    else:
        overview_rows.append(["失负荷率", _round(float(summary["loss_ratio_of_load"]), 6), "失负荷电量 / 总用电量"])
    overview_rows += [
        ["风电装机", _round(float(summary["wind_capacity_kw"]), 2), "kW"],
        ["光伏装机", _round(float(summary["pv_capacity_kw"]), 2), "kW"],
        ["储能容量", _round(float(summary["ess_energy_kwh"]), 2), "kWh"],
        ["储能功率", _round(float(summary["ess_power_kw"]), 2), "kW"],
    ]
    if diesel_enabled:
        overview_rows += [
            ["柴油装机", _round(float(summary["diesel_capacity_kw"]), 2), "kW"],
            ["柴油年发电量", _round(float(summary["total_diesel_generation_kwh"]), 2), "kWh/年"],
            ["柴油年燃料成本", _round(float(summary["annual_diesel_fuel_cost"]), 2), "元/年"],
        ]
    if biomass_enabled:
        overview_rows += [
            ["生物质装机", _round(float(summary["biomass_capacity_kw"]), 2), "kW"],
            ["生物质年发电量", _round(float(summary["total_biomass_generation_kwh"]), 2), "kWh/年"],
            ["生物质年燃料成本", _round(float(summary["annual_biomass_fuel_cost"]), 2), "元/年"],
        ]
    if maintenance_hours > 0:
        start_day = maintenance_start_hour // 24 + 1
        overview_rows.append([
            "生物质检修窗",
            f"第 {start_day} 天起，共 {maintenance_hours // 24} 天",
            "检修窗内出力为 0；时段由优化自动选择（通常在其他电源出力充裕处）",
        ])
    if solver_mode:
        overview_rows.append(["求解引擎模式", SOLVER_MODE_LABELS.get(solver_mode, solver_mode),
                              "双擎互证=两套独立引擎交叉核对一致后采信"])
    if boundary_notes:
        overview_rows.append(["边界提示", f"{len(boundary_notes)} 项", "详见「装机与边界」页"])
    _add_sheet(workbook, "方案总览", ["指标", "值", "说明"], overview_rows, widths=[26, 26, 58])

    # ── 经济性明细 ──
    econ_rows: list[list[Any]] = [
        ["运维成本", _round(float(summary["annual_om_cost"]), 2), _round(float(summary["om_cost_npv"]), 2)],
    ]
    if grid:
        econ_rows.append(["购电成本", _round(float(summary["annual_buy_cost"]), 2), _round(float(summary["buy_cost_npv"]), 2)])
        econ_rows.append(["售电收益（抵扣）", _round(float(summary["annual_sell_revenue"]), 2), _round(float(summary["sell_revenue_npv"]), 2)])
    else:
        econ_rows.append(["失负荷惩罚", _round(float(summary["annual_unmet_penalty_cost"]), 2), _round(float(summary["unmet_penalty_npv"]), 2)])
    if diesel_enabled:
        econ_rows.append(["柴油燃料成本", _round(float(summary["annual_diesel_fuel_cost"]), 2), _round(float(summary["diesel_fuel_cost_npv"]), 2)])
    if biomass_enabled:
        econ_rows.append(["生物质燃料成本", _round(float(summary["annual_biomass_fuel_cost"]), 2), _round(float(summary["biomass_fuel_cost_npv"]), 2)])
    econ_rows += [
        ["残值回收（抵扣）", "—", _round(float(summary["salvage_value_npv"]), 2)],
        ["折旧税盾（抵扣）", "—", _round(float(summary["tax_shield_npv"]), 2)],
    ]
    _add_sheet(workbook, "经济性明细", ["项目", "年值（元/年）", "现值（元）"], econ_rows, widths=[24, 20, 20])

    # ── 装机与边界 ──
    def _status(keyword: str) -> str:
        hits = [note for note in boundary_notes if keyword in note]
        return hits[0] if hits else "正常"

    bounds_rows: list[list[Any]] = []
    for label, key, summary_key in (
        ("风电", "wind", "wind_capacity_kw"),
        ("光伏", "pv", "pv_capacity_kw"),
        ("储能容量", "ess_energy", "ess_energy_kwh"),
        ("储能功率", "ess_power", "ess_power_kw"),
    ):
        lower, upper = bounds.get(key, (0.0, float("inf")))
        bounds_rows.append([
            label, _round(float(summary[summary_key]), 2),
            _round(lower, 2) if lower is not None else "—",
            _round(upper, 2) if upper is not None and upper != float("inf") else "不限",
            _status(label),
        ])
    if diesel_enabled:
        lower, upper = bounds.get("diesel", (0.0, float("inf")))
        bounds_rows.append([
            "柴油", _round(float(summary["diesel_capacity_kw"]), 2),
            _round(lower, 2), _round(upper, 2) if upper != float("inf") else "不限",
            _status("柴油装机"),
        ])
    if biomass_enabled:
        lower, upper = bounds.get("biomass", (0.0, float("inf")))
        bounds_rows.append([
            "生物质", _round(float(summary["biomass_capacity_kw"]), 2),
            _round(lower, 2), _round(upper, 2) if upper != float("inf") else "不限",
            _status("生物质装机"),
        ])
        if "biomass_energy" in bounds:
            energy_cap = bounds["biomass_energy"][1]
            bounds_rows.append([
                "生物质年发电量", _round(float(summary["total_biomass_generation_kwh"]), 2),
                0.0, _round(energy_cap, 2), _status("燃料年可用量上限"),
            ])
    for note in boundary_notes:
        bounds_rows.append(["提示", note, "", "", ""])
    _add_sheet(workbook, "装机与边界", ["组件", "结果值", "配置下限", "配置上限", "状态"],
               bounds_rows, widths=[18, 16, 14, 16, 46])

    # ── 逐时运行（按项目过滤列） ──
    skip_columns: set[str] = {"green_to_sell", "biomass_to_sell", "grid_to_load"} if not grid else {"unmet_load"}
    if not diesel_enabled:
        skip_columns.add("diesel_to_load")
    if not biomass_enabled:
        skip_columns |= {"biomass_generation", "biomass_to_load", "biomass_to_sell",
                         "biomass_curtail", "biomass_maintenance", "biomass_balance_check"}
    elif not biomass_continuous:
        skip_columns |= {"biomass_curtail", "biomass_maintenance"}
    hourly_columns = [
        name for name in HOURLY_COLUMN_LABELS
        if name in hourly_rows[0] and name not in skip_columns
    ]
    hourly_sheet = _add_sheet(
        workbook, "逐时运行",
        [HOURLY_COLUMN_LABELS[name] for name in hourly_columns],
        [[_round(row[name], 6) if isinstance(row[name], float) else row[name]
          for name in hourly_columns] for row in hourly_rows],
        widths=[12] + [16] * max(len(hourly_columns) - 1, 0),
    )
    hourly_sheet.freeze_panes = "A2"

    # ── 检修安排（仅持续运行检修启用时） ──
    if maintenance_hours > 0:
        start_day = maintenance_start_hour // 24 + 1
        wrap = maintenance_start_hour + maintenance_hours > len(hourly_rows)
        _add_sheet(workbook, "检修安排", ["项目", "内容"], [
            ["起始小时", maintenance_start_hour],
            ["起始日", f"第 {start_day} 天"],
            ["窗口时长", f"{maintenance_hours // 24} 天"],
            ["窗内出力策略", "强制为 0"],
            ["时段选择", "由优化自动选择，通常落在其他电源出力充裕、生物质电量价值最低的时段"],
            ["跨年衔接", "是（窗口跨越年末与次年年初）" if wrap else "否"],
        ], widths=[18, 60])

    excel_path = Path(excel_path)
    workbook.save(excel_path)
    return excel_path


def build_component_bounds(bounds: Any, diesel: Any, biomass: Any,
                           diesel_enabled: bool, biomass_enabled: bool) -> dict[str, tuple[float, float]]:
    """从配置节构造"装机与边界"页所需的组件 (下限, 上限) 表。

    bounds/diesel/biomass 为配置对象（鸭式类型，两套引擎共用）。仅启用
    组件进入结果；生物质年发电量上限仅在配置了燃料年可用量时出现。
    """
    result: dict[str, tuple[float, float]] = {
        "wind": (0.0, float(bounds.wind_capacity_kw_max)),
        "pv": (0.0, float(bounds.pv_capacity_kw_max)),
        "ess_energy": (0.0, float(bounds.ess_energy_kwh_max)),
        "ess_power": (0.0, float(bounds.ess_power_kw_max)),
    }
    if diesel_enabled:
        result["diesel"] = (float(diesel.min_capacity_kw), float(diesel.max_capacity_kw))
    if biomass_enabled:
        result["biomass"] = (float(biomass.min_capacity_kw), float(biomass.max_capacity_kw))
        if biomass.max_annual_energy_kwh is not None:
            result["biomass_energy"] = (0.0, float(biomass.max_annual_energy_kwh))
    return result


def append_validation_sheet(excel_path: str | Path, report: Any) -> Path:
    """将校验结论写入既有工作簿（汇总 + 仅未通过项）；工作簿不存在则创建。"""
    excel_path = Path(excel_path)
    if excel_path.exists():
        workbook = load_workbook(excel_path)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)
    if "校验结论" in workbook.sheetnames:
        workbook.remove(workbook["校验结论"])
    rows: list[list[Any]] = [
        ["校验结论", "通过" if report.passed else "未通过",
         f"{report.passed_checks}/{report.total_checks} 项检查通过"],
    ]
    if not report.passed:
        for check in report.checks:
            if not check.passed:
                rows.append([check.name, "未通过", check.details])
    _add_sheet(workbook, "校验结论", ["结果", "状态", "说明"], rows, widths=[40, 12, 70])
    workbook.save(excel_path)
    return excel_path
