"""收资清单读取器（v4.5）：智能体版收资清单 xlsx → 项目配置。

解析契约与清单模板「读取协议」Sheet 一致：
- 「项目与数据」表固定 7 行（按名称定位）；「参数」表 A 列为 config 字段路径；
- 填写值非空 = 业主已拍板 → 写入 config；留空 = 交智能体阶段3给推荐值；
- 必填缺失、路径不存在、疑似百分数、组件矛盾等一律进 issues（澄清项），
  不带猜测推进；已填但条件不符的参数进 ignored_params（附告警）。

本模块保持纯解析（不做文件写入）；项目配置的落盘与校验由 agent 层完成。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROFILE_SHEET = "项目与数据"
PARAM_SHEET = "参数"
PROFILE_FIELDS = ("项目名称", "项目类型", "纳入柴油发电", "纳入生物质发电",
                  "负荷数据文件路径", "电价文件路径", "备注")
PROFILE_PARAMS = {"project.type", "diesel.enabled", "biomass.enabled"}
YEAR_PATHS = {"finance.lifecycle_years", "finance.depreciation_years"}
# 比例/效率类字段：必须 ≤1（疑似百分数检测）；储能时长为小时数，不受此限
RATIO_PATHS_PREFIX = ("policy.", "biomass.min_output_ratio", "biomass.capacity_factor",
                      "project.max_loss_ratio_of_load")
SOLVER_MODE_ALIASES = {"default": "default", "alternate": "alternate", "dual": "dual",
                       "天枢": "default", "天璇": "alternate", "双擎互证": "dual"}
PROJECT_TYPE_ALIASES = {"并网型": "grid_connected", "离网型": "offgrid",
                        "grid_connected": "grid_connected", "offgrid": "offgrid"}
YES_ALIASES = {"是", "true", "y", "yes"}
NO_ALIASES = {"否", "false", "n", "no", ""}


@dataclass(frozen=True)
class IntakeResult:
    project_name: str
    project_type: str
    diesel_enabled: bool
    biomass_enabled: bool
    load_path: Path | None
    price_path: Path | None
    config_overrides: dict
    filled_params: tuple[tuple[str, Any], ...]
    blank_params: tuple[str, ...]
    ignored_params: tuple[str, ...]
    issues: tuple[str, ...]
    warnings: tuple[str, ...]
    notes: str | None

    @property
    def ok(self) -> bool:
        return not self.issues


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _as_number(path: str, value: Any, row: int, issues: list[str]) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        issues.append(f"「{PARAM_SHEET}」表第 {row} 行 {path}：填写值「{value}」无法识别为数字，请填写数值。")
        return None


def _applicable(path: str, project_type: str, diesel_enabled: bool, biomass_enabled: bool) -> bool:
    if path == "project.loss_of_load_penalty_per_kwh" or path == "project.max_loss_ratio_of_load":
        return project_type == "offgrid"
    if path in ("policy.max_sell_ratio_of_green", "bounds.grid_import_kw_max", "bounds.grid_export_kw_max"):
        return project_type == "grid_connected"
    if path.startswith("diesel."):
        return diesel_enabled
    if path.startswith("biomass."):
        return biomass_enabled
    return True


def _set_path(target: dict, dotted: str, value: Any) -> None:
    parts = dotted.split(".")
    node = target
    for part in parts[:-1]:
        node = node.setdefault(part, {})
    node[parts[-1]] = value


def parse_intake_checklist(checklist_path: str | Path) -> IntakeResult:
    issues: list[str] = []
    warnings: list[str] = []
    path = Path(checklist_path).expanduser()
    if not path.is_file():
        return IntakeResult("", "grid_connected", False, False, None, None, {}, (), (), (),
                            (f"收资清单文件不存在：{path}",), (), None)

    workbook = load_workbook(path, data_only=True)
    for required_sheet in (PROFILE_SHEET, PARAM_SHEET):
        if required_sheet not in workbook.sheetnames:
            return IntakeResult("", "grid_connected", False, False, None, None, {}, (), (), (),
                                (f"这不是智能体版收资清单：缺少「{required_sheet}」表，"
                                 "请使用《绿电直连项目收资清单_智能体版》模板填写。",), (), None)

    # ── 项目与数据 ───────────────────────────────────────────────
    profile = {PROFILE_SHEET: {}}
    profile_rows: dict[str, tuple[Any, int]] = {}
    for row in range(2, workbook[PROFILE_SHEET].max_row + 1):
        name = workbook[PROFILE_SHEET].cell(row=row, column=1).value
        if name:
            profile_rows[str(name).strip()] = (workbook[PROFILE_SHEET].cell(row=row, column=2).value, row)
    for field_name in PROFILE_FIELDS:
        if field_name not in profile_rows:
            issues.append(f"「{PROFILE_SHEET}」表缺少「{field_name}」行，请使用原模板，不要删行。")
    if issues:
        return IntakeResult("", "grid_connected", False, False, None, None, {}, (), (),
                            tuple(issues), (), None)

    def profile_value(field_name: str) -> Any:
        return profile_rows[field_name][0]

    def profile_row(field_name: str) -> int:
        return profile_rows[field_name][1]

    name_raw = profile_value("项目名称")
    if _is_blank(name_raw):
        issues.append(f"「{PROFILE_SHEET}」表：项目名称未填写（必填，将用作归档目录名）。")
        project_name = ""
    else:
        project_name = str(name_raw).strip()
        if any(ch in project_name for ch in '/\\:'):
            issues.append(f"「{PROFILE_SHEET}」表：项目名称「{project_name}」含路径非法字符（/ \\ :），请修改。")

    type_raw = str(profile_value("项目类型") or "").strip()
    if type_raw.lower() in PROJECT_TYPE_ALIASES:
        project_type = PROJECT_TYPE_ALIASES[type_raw.lower()]
    elif _is_blank(type_raw):
        issues.append(f"「{PROFILE_SHEET}」表：项目类型未填写（下拉二选一：并网型 / 离网型）。")
        project_type = "grid_connected"
    else:
        issues.append(f"「{PROFILE_SHEET}」表：项目类型「{type_raw}」无法识别，请从下拉选择并网型 / 离网型。")
        project_type = "grid_connected"

    def yes_no(field_name: str) -> bool:
        raw = str(profile_value(field_name) or "").strip().lower()
        if raw in YES_ALIASES:
            return True
        if raw in NO_ALIASES:
            if _is_blank(profile_value(field_name)):
                warnings.append(f"「{PROFILE_SHEET}」表「{field_name}」未填写，按「否」处理。")
            return False
        issues.append(f"「{PROFILE_SHEET}」表：「{field_name}」填写值「{profile_value(field_name)}」无法识别，请填 是 / 否。")
        return False

    diesel_enabled = yes_no("纳入柴油发电")
    biomass_enabled = yes_no("纳入生物质发电")

    def data_file(field_name: str, required: bool) -> Path | None:
        raw = profile_value(field_name)
        if _is_blank(raw):
            if required:
                issues.append(f"「{PROFILE_SHEET}」表：{field_name}未填写（必填，本机绝对路径）。")
            return None
        file_path = Path(str(raw).strip()).expanduser()
        if not file_path.exists():
            issues.append(f"「{PROFILE_SHEET}」表：{field_name}指向的文件不存在：{file_path}")
        elif file_path.is_dir():
            issues.append(f"「{PROFILE_SHEET}」表：{field_name}指向的是目录，请填写具体数据文件（csv/xlsx/xls）。")
        return file_path

    load_path = data_file("负荷数据文件路径", required=True)
    price_path = data_file("电价文件路径", required=(project_type == "grid_connected"))
    if project_type == "grid_connected" and price_path is None and not any("电价" in text for text in issues):
        pass  # required 分支已生成 issue
    notes = None if _is_blank(profile_value("备注")) else str(profile_value("备注")).strip()

    # ── 参数表 ───────────────────────────────────────────────────
    sheet = workbook[PARAM_SHEET]
    filled: list[tuple[str, Any, int]] = []
    blank: list[str] = []
    ignored: list[str] = []
    overrides: dict = {}
    for row in range(2, sheet.max_row + 1):
        path_value = sheet.cell(row=row, column=1).value
        if _is_blank(path_value):
            continue
        dotted = str(path_value).strip()
        cell_value = sheet.cell(row=row, column=2).value
        if dotted == "solver.mode":
            if _is_blank(cell_value):
                continue  # 留空 = 天枢（默认），不需要智能体再征询
            mode = SOLVER_MODE_ALIASES.get(str(cell_value).strip())
            if mode is None:
                issues.append(f"「{PARAM_SHEET}」表第 {row} 行 solver.mode：填写值「{cell_value}」无法识别，"
                              "请填 default / alternate / dual。")
            else:
                _set_path(overrides, "solver.mode", mode)
                filled.append((dotted, mode, row))
            continue
        if _is_blank(cell_value):
            if _applicable(dotted, project_type, diesel_enabled, biomass_enabled):
                blank.append(dotted)
            continue
        if not _applicable(dotted, project_type, diesel_enabled, biomass_enabled):
            ignored.append(dotted)
            warnings.append(f"「{PARAM_SHEET}」表第 {row} 行 {dotted}：与项目条件不符（未纳入对应组件或项目类型不同），已忽略。")
            continue
        number = _as_number(dotted, cell_value, row, issues)
        if number is None:
            continue
        if dotted in YEAR_PATHS:
            if not number.is_integer():
                issues.append(f"「{PARAM_SHEET}」表第 {row} 行 {dotted}：年限须为整数，当前为 {number}。")
                continue
            value: Any = int(number)
        else:
            value = number
        if (dotted.startswith(RATIO_PATHS_PREFIX) or "ratio" in dotted or "efficiency" in dotted) and value > 1.0001:
            issues.append(f"「{PARAM_SHEET}」表第 {row} 行 {dotted}：填写值 {value} 大于 1，"
                          "疑似填了百分数——比例请填小数（如 8% 填 0.08）。")
            continue
        if value < 0:
            issues.append(f"「{PARAM_SHEET}」表第 {row} 行 {dotted}：填写值不能为负数。")
            continue
        _set_path(overrides, dotted, value)
        filled.append((dotted, value, row))

    # 并网项目：负荷文件同时提供出力列即可（体检核对）；离网项目电价留空合法。
    overrides.setdefault("project", {})["type"] = project_type
    if diesel_enabled:
        _set_path(overrides, "diesel.enabled", True)
    if biomass_enabled:
        _set_path(overrides, "biomass.enabled", True)

    blank_params = tuple(path for path in blank if path not in {item[0] for item in filled})
    return IntakeResult(
        project_name=project_name,
        project_type=project_type,
        diesel_enabled=diesel_enabled,
        biomass_enabled=biomass_enabled,
        load_path=load_path,
        price_path=price_path,
        config_overrides=overrides,
        filled_params=tuple((item[0], item[1]) for item in filled),
        blank_params=blank_params,
        ignored_params=tuple(ignored),
        issues=tuple(issues),
        warnings=tuple(warnings),
        notes=notes,
    )
