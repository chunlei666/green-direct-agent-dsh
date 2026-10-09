"""业主交付报告（v4.4）：Word 文档，由引擎确定性生成。

设计约定：
- 报告全部数字直接取自求解结果与配置，大脑不撰写、不改写（结论摘要
  以载荷 summary_points 逐字转述）。
- 结构固定：项目简介 / 结论摘要 / 推荐方案 / 经济性 / 运行特性 / 附录。
- 按项目组件过滤：禁用组件（柴油/生物质）通篇不出现。
- 文风：结论先行、每条一句；解读句不超过两句；无填充套话。
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from .config import OptimizationConfig
from .excel_report import build_component_bounds

_ZERO_CAPACITY_TOL = 1.0e-6


def _num(value: Any) -> float:
    return float(value)


def _wan(value: float) -> str:
    return f"{value / 1.0e4:,.1f}"


def _pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def _kw(value: float) -> str:
    return f"{round(value, 1):,.10g}"


def build_summary_points(summary: dict[str, Any], config: OptimizationConfig) -> list[str]:
    """结论摘要（≤4 条，每条一句）；报告正文与载荷共用同一来源。"""
    grid = config.project.type == "grid_connected"
    parts: list[str] = []
    if _num(summary["wind_capacity_kw"]) > _ZERO_CAPACITY_TOL:
        parts.append(f"风电 {_kw(_num(summary['wind_capacity_kw']))} kW")
    if _num(summary["pv_capacity_kw"]) > _ZERO_CAPACITY_TOL:
        parts.append(f"光伏 {_kw(_num(summary['pv_capacity_kw']))} kW")
    if _num(summary["ess_energy_kwh"]) > _ZERO_CAPACITY_TOL or _num(summary["ess_power_kw"]) > _ZERO_CAPACITY_TOL:
        parts.append(
            f"储能 {_kw(_num(summary['ess_energy_kwh']))} kWh / {_kw(_num(summary['ess_power_kw']))} kW"
        )
    if config.diesel.enabled and _num(summary["diesel_capacity_kw"]) > _ZERO_CAPACITY_TOL:
        parts.append(f"柴油 {_kw(_num(summary['diesel_capacity_kw']))} kW")
    if config.biomass.enabled and _num(summary["biomass_capacity_kw"]) > _ZERO_CAPACITY_TOL:
        parts.append(f"生物质 {_kw(_num(summary['biomass_capacity_kw']))} kW")
    if parts:
        first = "推荐配置：" + "、".join(parts) + "。"
    else:
        first = "推荐配置：经优化测算，各候选组件均不宜配置（详见推荐方案节）。"

    points = [
        first,
        f"全生命周期总成本 {_wan(_num(summary['objective_value']))} 万元，"
        f"其中初始投资 {_wan(_num(summary['initial_investment_cost']))} 万元。",
    ]
    if grid:
        points.append(
            f"关键指标：绿电自用率 {_pct(_num(summary['self_use_ratio_of_green']))}，"
            f"弃电率 {_pct(_num(summary['curtail_ratio_of_green']))}。"
        )
    else:
        points.append(
            f"关键指标：绿电自用率 {_pct(_num(summary['self_use_ratio_of_green']))}，"
            f"失负荷率 {_pct(_num(summary['loss_ratio_of_load']))}。"
        )
    maint_hours = int(round(config.biomass.maintenance_days * 24)) if config.biomass.enabled else 0
    if maint_hours > 0 and _num(summary.get("biomass_maintenance_hours", 0)) > 0:
        start_hour = int(_num(summary["biomass_maintenance_start_hour"]))
        points.append(
            f"检修安排：第 {start_hour // 24 + 1} 天起共 {maint_hours // 24} 天（优化自动选定）。"
        )
    return points


def _zero_reason(name: str, config: OptimizationConfig, summary: dict[str, Any]) -> str:
    """启用但装机为 0 的组件，给一句话经济归因。"""
    grid = config.project.type == "grid_connected"
    if name == "biomass" and grid:
        fuel = config.biomass.fuel_cost_per_kwh
        purchase_kwh = _num(summary.get("total_grid_purchase_kwh", 0.0))
        if purchase_kwh > _ZERO_CAPACITY_TOL:
            avg_buy = _num(summary["total_buy_cost"]) / purchase_kwh
            return (
                f"装机为 0：生物质度电燃料成本（{fuel:g} 元/kWh）相对平均购电价"
                f"（{avg_buy:.3f} 元/kWh）无优势，叠加投资后不经济，经优化不建设。"
            )
        return "装机为 0：经全生命周期经济测算，配置生物质不占优，故不建设。"
    if name == "biomass":
        return "装机为 0：经全生命周期经济测算，配置生物质不占优，故不建设（如需保留可设置最小容量后重新求解）。"
    if name in ("wind", "pv") and grid:
        return f"装机为 0：其单位电量综合成本高于购电价，经优化不建设。"
    labels = {"wind": "风电", "pv": "光伏", "ess": "储能", "diesel": "柴油"}
    return f"装机为 0：经全生命周期经济测算，配置{labels.get(name, '该组件')}不占优，故不建设。"


def _add_table(document: Document, headers: list[str], rows: list[list[Any]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            cells[index].text = str(text)
    document.add_paragraph()


def _econ_rows(summary: dict[str, Any], config: OptimizationConfig) -> list[list[Any]]:
    grid = config.project.type == "grid_connected"
    rows: list[list[Any]] = [["初始投资", "—", _wan(_num(summary["initial_investment_cost"]))]]
    rows.append(["运维成本", _wan(_num(summary["annual_om_cost"])), _wan(_num(summary["om_cost_npv"]))])
    if grid:
        rows.append(["购电成本", _wan(_num(summary["annual_buy_cost"])), _wan(_num(summary["buy_cost_npv"]))])
        rows.append(["售电收益（抵扣）", _wan(_num(summary["annual_sell_revenue"])), _wan(_num(summary["sell_revenue_npv"]))])
    else:
        rows.append(["失负荷惩罚", _wan(_num(summary["annual_unmet_penalty_cost"])), _wan(_num(summary["unmet_penalty_npv"]))])
    if config.diesel.enabled:
        rows.append(["柴油燃料成本", _wan(_num(summary["annual_diesel_fuel_cost"])), _wan(_num(summary["diesel_fuel_cost_npv"]))])
    if config.biomass.enabled:
        rows.append(["生物质燃料成本", _wan(_num(summary["annual_biomass_fuel_cost"])), _wan(_num(summary["biomass_fuel_cost_npv"]))])
    rows.append(["残值回收（抵扣）", "—", _wan(_num(summary["salvage_value_npv"]))])
    rows.append(["折旧税盾（抵扣）", "—", _wan(_num(summary["tax_shield_npv"]))])
    return rows


def _cost_leader(summary: dict[str, Any], config: OptimizationConfig) -> str:
    """一句话点出成本大头（按现值最大的单项成本）。"""
    grid = config.project.type == "grid_connected"
    candidates = [("初始投资", _num(summary["initial_investment_cost"]))]
    if _num(summary["om_cost_npv"]) > 0:
        candidates.append(("运维成本", _num(summary["om_cost_npv"])))
    if grid and _num(summary["buy_cost_npv"]) > 0:
        candidates.append(("购电成本", _num(summary["buy_cost_npv"])))
    if not grid and _num(summary["unmet_penalty_npv"]) > 0:
        candidates.append(("失负荷惩罚", _num(summary["unmet_penalty_npv"])))
    if config.diesel.enabled and _num(summary["diesel_fuel_cost_npv"]) > 0:
        candidates.append(("柴油燃料成本", _num(summary["diesel_fuel_cost_npv"])))
    if config.biomass.enabled and _num(summary["biomass_fuel_cost_npv"]) > 0:
        candidates.append(("生物质燃料成本", _num(summary["biomass_fuel_cost_npv"])))
    leader, value = max(candidates, key=lambda item: item[1])
    total = _num(summary["objective_value"])
    share = f"，约占全生命周期总成本的 {_pct(value / total)}" if total > 1.0e-9 else ""
    return f"成本大头为{leader}（现值 {_wan(value)} 万元{share}）。"


def _appendix_rows(config: OptimizationConfig) -> list[list[Any]]:
    f = config.finance
    rows: list[list[Any]] = [
        ["时间步长", f"{config.timestep_hours:g}", "小时"],
        ["折现率", f"{f.discount_rate:.2%}", "—"],
        ["全生命周期", f"{f.lifecycle_years}", "年"],
        ["折旧年限", f"{f.depreciation_years}", "年"],
        ["所得税率", f"{f.income_tax_rate:.0%}", "—"],
        ["残值率", f"{f.salvage_rate:.0%}", "—"],
    ]
    grid = config.project.type == "grid_connected"
    if grid:
        rows.append(["购电价 / 售电价", "按逐时电价文件执行", "—"])
    else:
        rows.append(["失负荷惩罚单价", f"{config.project.loss_of_load_penalty_per_kwh:g}", "元/kWh"])
        cap = config.project.max_loss_ratio_of_load
        rows.append(["最大失负荷率", "不限" if cap is None else f"{cap:.1%}", "—"])
    b = config.bounds
    rows += [
        ["风电造价 / 运维", f"{config.costs.wind_capex_per_kw:g} 元/kW；{config.costs.wind_om_per_kw_year:g} 元/kW·年；装机 0~{_kw(b.wind_capacity_kw_max)} kW", "—"],
        ["光伏造价 / 运维", f"{config.costs.pv_capex_per_kw:g} 元/kW；{config.costs.pv_om_per_kw_year:g} 元/kW·年；装机 0~{_kw(b.pv_capacity_kw_max)} kW", "—"],
        ["储能造价 / 效率 / 约束", (
            f"能量 {config.costs.ess_energy_capex_per_kwh:g} 元/kWh；功率 {config.costs.ess_power_capex_per_kw:g} 元/kW；"
            f"充/放效率 {config.storage.charge_efficiency:.0%}/{config.storage.discharge_efficiency:.0%}；"
            f"SOC {config.storage.soc_min_ratio:.0%}~{config.storage.soc_max_ratio:.0%}；"
            f"时长 {config.storage.min_storage_hours:g}~{config.storage.max_storage_hours:g} 小时；"
            f"容量 0~{_kw(b.ess_energy_kwh_max)} kWh / {_kw(b.ess_power_kw_max)} kW"
        ), "—"],
    ]
    if config.diesel.enabled:
        d = config.diesel
        rows.append(["柴油参数", (
            f"造价 {d.capex_per_kw:g} 元/kW；运维 {d.fixed_om_per_kw_year:g} 元/kW·年；"
            f"燃料 {d.fuel_cost_per_kwh:g} 元/kWh；装机 {d.min_capacity_kw:g}~{d.max_capacity_kw:g} kW"
        ), "—"])
    if config.biomass.enabled:
        m = config.biomass
        fuel_cap = "不限" if m.max_annual_energy_kwh is None else f"{m.max_annual_energy_kwh:,.0f} kWh"
        rows.append(["生物质参数", (
            f"造价 {m.capex_per_kw:g} 元/kW；运维 {m.fixed_om_per_kw_year:g} 元/kW·年；"
            f"燃料 {m.fuel_cost_per_kwh:g} 元/kWh；容量因子 {m.capacity_factor:.0%}；"
            f"装机 {m.min_capacity_kw:g}~{m.max_capacity_kw:g} kW；燃料年上限 {fuel_cap}；"
            f"最小出力比例 {m.min_output_ratio:g}；年检修 {m.maintenance_days:g} 天"
        ), "—"])
    p = config.policy
    policy_text = f"绿电自用率下限 {p.min_self_use_ratio_of_green:.0%}；负荷绿电覆盖率下限 {p.min_self_use_ratio_of_load:.0%}"
    if grid:
        policy_text += f"；绿电外售率上限 {p.max_sell_ratio_of_green:.0%}"
    rows.append(["政策参数", policy_text, "—"])
    return rows


def write_report_docx(
    output_path: str | Path,
    *,
    summary: dict[str, Any],
    config: OptimizationConfig,
    project_name: str,
) -> tuple[Path, list[str]]:
    """生成业主交付 Word 报告；返回（文件路径，结论摘要文本列表）。"""
    grid = config.project.type == "grid_connected"
    points = build_summary_points(summary, config)
    bounds = build_component_bounds(
        config.bounds, config.diesel, config.biomass,
        config.diesel.enabled, config.biomass.enabled,
    )

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    document.add_heading("绿电直连项目配置优化报告", 0)
    subtitle = document.add_paragraph(f"（{project_name} · {date.today().isoformat()}）")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("一、项目简介", level=1)
    components = ["风电", "光伏", "储能"]
    if config.diesel.enabled:
        components.append("柴油")
    if config.biomass.enabled:
        components.append("生物质")
    document.add_paragraph(
        f"本项目为{PROJECT_TYPE_LABEL}（{PROJECT_TYPE_TEXT[grid]}），"
        f"拟通过{'、'.join(components)}的一体化配置实现绿电高效供给与经济运行。"
        f"本报告基于项目方提供的负荷与资源时序数据、价格与财务参数，"
        f"对各电源与储能配置方案进行全生命周期经济优化，为投资决策提供依据。"
    )
    document.add_paragraph("项目地点：【待补充】；项目业主：【待补充】；编制说明：本报告由优化引擎依据求解结果自动生成。")

    document.add_heading("二、结论摘要", level=1)
    for index, point in enumerate(points, start=1):
        document.add_paragraph(f"{index}. {point}")

    document.add_heading("三、推荐方案", level=1)
    plan_rows: list[list[Any]] = []
    for label, key, bound_key, unit in (
        ("风电", "wind_capacity_kw", "wind", "kW"),
        ("光伏", "pv_capacity_kw", "pv", "kW"),
        ("储能容量", "ess_energy_kwh", "ess_energy", "kWh"),
        ("储能功率", "ess_power_kw", "ess_power", "kW"),
    ):
        lower, upper = bounds[bound_key]
        plan_rows.append([label, f"{_kw(_num(summary[key]))} {unit}", f"{_kw(lower)} ~ {_kw(upper)} {unit}"])
    if config.diesel.enabled:
        lower, upper = bounds["diesel"]
        plan_rows.append(["柴油", f"{_kw(_num(summary['diesel_capacity_kw']))} kW", f"{_kw(lower)} ~ {_kw(upper)} kW"])
    if config.biomass.enabled:
        lower, upper = bounds["biomass"]
        plan_rows.append(["生物质", f"{_kw(_num(summary['biomass_capacity_kw']))} kW", f"{_kw(lower)} ~ {_kw(upper)} kW"])
    _add_table(document, ["组件", "装机", "配置范围"], plan_rows)
    if config.biomass.enabled and _num(summary["biomass_capacity_kw"]) <= _ZERO_CAPACITY_TOL:
        document.add_paragraph(_zero_reason("biomass", config, summary))
    if config.diesel.enabled and _num(summary["diesel_capacity_kw"]) <= _ZERO_CAPACITY_TOL:
        document.add_paragraph(_zero_reason("diesel", config, summary))
    if _num(summary["wind_capacity_kw"]) <= _ZERO_CAPACITY_TOL:
        document.add_paragraph(_zero_reason("wind", config, summary))
    if _num(summary["pv_capacity_kw"]) <= _ZERO_CAPACITY_TOL:
        document.add_paragraph(_zero_reason("pv", config, summary))

    document.add_heading("四、经济性", level=1)
    _add_table(document, ["成本项", "年值（万元/年）", "现值（万元）"], _econ_rows(summary, config))
    document.add_paragraph(_cost_leader(summary, config))

    document.add_heading("五、运行特性", level=1)
    if grid:
        document.add_paragraph(
            f"项目绿电自用率 {_pct(_num(summary['self_use_ratio_of_green']))}，"
            f"弃电率 {_pct(_num(summary['curtail_ratio_of_green']))}。"
        )
    else:
        document.add_paragraph(
            f"项目绿电自用率 {_pct(_num(summary['self_use_ratio_of_green']))}，"
            f"失负荷率 {_pct(_num(summary['loss_ratio_of_load']))}。"
        )
    if config.biomass.enabled:
        text = f"生物质年发电量 {_num(summary['total_biomass_generation_kwh']):,.0f} kWh"
        if "biomass_energy" in bounds:
            cap = bounds["biomass_energy"][1]
            share = _num(summary["total_biomass_generation_kwh"]) / cap if cap > 0 else 0.0
            text += f"，为燃料年可用量上限（{cap:,.0f} kWh）的 {_pct(share)}"
        document.add_paragraph(text + "。")
    maint_hours = int(round(config.biomass.maintenance_days * 24)) if config.biomass.enabled else 0
    if maint_hours > 0 and _num(summary.get("biomass_maintenance_hours", 0)) > 0:
        start_hour = int(_num(summary["biomass_maintenance_start_hour"]))
        document.add_paragraph(
            f"检修窗安排在第 {start_hour // 24 + 1} 天起共 {maint_hours // 24} 天，窗内机组停机、出力为 0。"
            "该时段由优化在全年自动选定，通常为其他电源出力充裕、生物质电量价值最低的时段，对供电与经济性影响最小。"
        )
    document.add_paragraph("逐时运行明细详见成果工作簿「逐时运行」页。")

    document.add_heading("六、附录", level=1)
    document.add_paragraph("1. 数据口径（本项目参数汇总）")
    _add_table(document, ["参数", "取值", "单位"], _appendix_rows(config))
    document.add_paragraph("2. 交付文件清单")
    document.add_paragraph(
        "本报告（green_direct_report.docx）；成果工作簿（green_direct_results.xlsx，"
        "含方案总览、经济性明细、装机与边界、逐时运行等页，持续运行项目另有检修安排页，"
        "校验后含校验结论页）。"
    )

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path, points


PROJECT_TYPE_LABEL = "绿电直连项目"
PROJECT_TYPE_TEXT = {True: "并网型", False: "离网型"}
